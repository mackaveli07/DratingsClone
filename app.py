# storage.py
import os
import time
from contextlib import contextmanager

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from config import EXCEL_FILE, HIST_SHEET, SCHEDULE_SHEET
from team_utils import map_team_name

try:
    import fcntl
except ImportError:
    fcntl = None


def load_games(file=EXCEL_FILE):
    if os.path.exists(file):
        try:
            hist_df = pd.read_excel(file, sheet_name=HIST_SHEET)
        except Exception:
            hist_df = pd.DataFrame()
        try:
            sched_df = pd.read_excel(file, sheet_name=SCHEDULE_SHEET)
        except Exception:
            sched_df = pd.DataFrame()
        return hist_df, sched_df
    return pd.DataFrame(), pd.DataFrame()


@pd.api.extensions.register_dataframe_accessor("saved_picks")
def load_saved_picks(file=EXCEL_FILE):
    columns = ["week", "matchup", "pick", "timestamp"]
    if not os.path.exists(file):
        return pd.DataFrame(columns=columns)
    try:
        df = pd.read_excel(file, sheet_name="Picks")
    except ValueError as exc:
        if "Worksheet named 'Picks' not found" in str(exc):
            return pd.DataFrame(columns=columns)
        raise
    for col in columns:
        if col not in df.columns:
            df[col] = np.nan
    return df[columns]


def _read_saved_picks_from_excel(file=EXCEL_FILE):
    return load_saved_picks(file)


def _write_picks_sheet(file, picks_df, columns):
    if os.path.exists(file):
        workbook = load_workbook(file)
        if "Picks" in workbook.sheetnames:
            sheet = workbook["Picks"]
        else:
            sheet = workbook.create_sheet("Picks")
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Picks"

    if sheet.max_row and sheet.max_row > 0:
        sheet.delete_rows(1, sheet.max_row)
    for col_idx, col_name in enumerate(columns, start=1):
        sheet.cell(row=1, column=col_idx, value=col_name)
    for row_idx, row in enumerate(picks_df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=value)
    workbook.save(file)


@contextmanager
def picks_file_lock(file):
    lock_path = f"{file}.lock"
    if fcntl:
        with open(lock_path, "w") as lock_file:
            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)
        return

    lock_dir = f"{lock_path}.d"
    acquired = False
    try:
        for _ in range(100):
            try:
                os.mkdir(lock_dir)
                acquired = True
                break
            except FileExistsError:
                time.sleep(0.05)
        if not acquired:
            raise TimeoutError("Could not acquire picks save lock.")
        yield
    finally:
        if acquired and os.path.isdir(lock_dir):
            os.rmdir(lock_dir)


def save_week_picks(week, picks_dict, file=EXCEL_FILE):
    columns = ["week", "matchup", "pick", "timestamp"]
    try:
        week_int = int(week)
    except Exception:
        return False

    rows = []
    now_ts = pd.Timestamp.now().strftime("%Y-%m-%dT%H:%M:%S")
    for matchup, pick in (picks_dict or {}).items():
        norm_matchup = matchup
        norm_pick = map_team_name(pick)
        if not norm_matchup or not norm_pick:
            continue
        rows.append({"week": week_int, "matchup": norm_matchup, "pick": norm_pick, "timestamp": now_ts})

    if not rows:
        return False

    new_rows = pd.DataFrame(rows, columns=columns)

    def _save_once():
        existing = _read_saved_picks_from_excel(file).copy()
        if not existing.empty:
            existing["week"] = pd.to_numeric(existing["week"], errors="coerce")
            existing["pick"] = existing["pick"].apply(map_team_name)
            existing = existing.dropna(subset=["week", "matchup", "pick"])
            existing["week"] = existing["week"].astype(int)
            existing = existing[~((existing["week"] == week_int) & (existing["matchup"].isin(new_rows["matchup"])))]

        out = pd.concat([existing[columns], new_rows], ignore_index=True)
        out = out.drop_duplicates(subset=["week", "matchup"], keep="last")
        _write_picks_sheet(file, out, columns)

    with picks_file_lock(file):
        _save_once()

    return True


def build_actual_results_by_week(hist_df):
    needed = {"week", "team1", "team2", "score1", "score2"}
    cols = ["week", "matchup", "winner", "is_final"]
    if hist_df is None or hist_df.empty or not needed.issubset(set(hist_df.columns)):
        return pd.DataFrame(columns=cols)

    rows = []
    for _, row in hist_df.iterrows():
        week = pd.to_numeric(row.get("week"), errors="coerce")
        if pd.isna(week):
            continue
        away_team = map_team_name(row.get("team1"))
        home_team = map_team_name(row.get("team2"))
        score1 = pd.to_numeric(row.get("score1"), errors="coerce")
        score2 = pd.to_numeric(row.get("score2"), errors="coerce")
        score_complete = pd.notna(score1) and pd.notna(score2)

        status_raw = row.get("status", row.get("game_status", row.get("state", None)))
        status_text = str(status_raw).strip().lower() if pd.notna(status_raw) else ""
        normalized_status = " ".join(status_text.replace("-", " ").replace("/", " ").split())
        status_tokens = set(normalized_status.split()) if normalized_status else set()

        if "postponed" in status_tokens:
            is_final = False
        elif "final" in status_tokens or {"complete", "completed"} & status_tokens or normalized_status == "post":
            is_final = True
        elif (
            normalized_status.startswith("q")
            or "live" in status_tokens
            or "halftime" in status_tokens
            or "scheduled" in status_tokens
            or "pregame" in status_tokens
            or "pre" in status_tokens
            or ("in" in status_tokens and "progress" in status_tokens)
        ):
            is_final = False
        else:
            is_final = score_complete

        winner = None
        if is_final:
            if score1 > score2:
                winner = away_team
            elif score2 > score1:
                winner = home_team
            else:
                winner = "TIE"

        rows.append({
            "week": int(week),
            "matchup": f"{away_team} @ {home_team}",
            "winner": winner,
            "is_final": bool(is_final),
        })

    if not rows:
        return pd.DataFrame(columns=cols)
    return pd.DataFrame(rows, columns=cols).drop_duplicates(subset=["week", "matchup"], keep="last")


def grade_picks(saved_picks_df, results_df):
    graded_cols = ["week", "matchup", "pick", "timestamp", "winner", "is_final", "status", "result"]
    if saved_picks_df is None or saved_picks_df.empty:
        return pd.DataFrame(columns=graded_cols)

    picks = saved_picks_df.copy()
    for col in ["week", "matchup", "pick", "timestamp"]:
        if col not in picks.columns:
            picks[col] = np.nan

    picks["week"] = pd.to_numeric(picks["week"], errors="coerce")
    picks = picks.dropna(subset=["week", "matchup", "pick"])
    if picks.empty:
        return pd.DataFrame(columns=graded_cols)

    picks["week"] = picks["week"].astype(int)
    picks["pick"] = picks["pick"].apply(map_team_name)

    if results_df is None or results_df.empty:
        merged = picks.copy()
        merged["winner"] = None
        merged["is_final"] = False
    else:
        results = results_df.copy()
        results["week"] = pd.to_numeric(results["week"], errors="coerce")
        results = results.dropna(subset=["week", "matchup"])
        results["week"] = results["week"].astype(int)
        merged = picks.merge(results[["week", "matchup", "winner", "is_final"]], on=["week", "matchup"], how="left")
        merged["is_final"] = merged["is_final"].fillna(False)

    merged["status"] = np.where(
        merged["is_final"] != True,
        "pending",
        np.where(
            merged["winner"] == "TIE",
            "tie/push",
            np.where(merged["pick"] == merged["winner"], "correct", "wrong"),
        ),
    )
    merged["result"] = np.where(
        merged["status"] == "correct",
        "✅ Correct",
        np.where(
            merged["status"] == "wrong",
            "❌ Wrong",
            np.where(merged["status"] == "tie/push", "➖ Tie/Push", "⏳ Pending"),
        ),
    )
    return merged[graded_cols]


__all__ = [
    "load_games",
    "load_saved_picks",
    "_read_saved_picks_from_excel",
    "_write_picks_sheet",
    "picks_file_lock",
    "save_week_picks",
    "build_actual_results_by_week",
    "grade_picks",
]

