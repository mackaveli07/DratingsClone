import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook

import app
from config import HIST_SHEET, PICKS_SHEET


class StorageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)

    def test_load_saved_picks_missing_sheet_returns_empty_frame(self) -> None:
        workbook_path = self.root / "games.xlsx"
        workbook = Workbook()
        workbook.active.title = HIST_SHEET
        workbook.save(workbook_path)

        picks = app.load_saved_picks(workbook_path)

        self.assertTrue(picks.empty)
        self.assertEqual(list(picks.columns), app.PICKS_COLUMNS)

    def test_load_games_preserves_missing_sheet_behavior(self) -> None:
        workbook_path = self.root / "games.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = HIST_SHEET
        sheet.append(["week", "team1", "team2", "score1", "score2"])
        sheet.append([1, "BUF", "NYJ", 24, 21])
        workbook.save(workbook_path)

        hist_df, sched_df = app.load_games(workbook_path)

        self.assertEqual(hist_df.iloc[0]["team1"], "BUF")
        self.assertTrue(sched_df.empty)

    def test_save_week_picks_rejects_malformed_matchup(self) -> None:
        workbook_path = self.root / "games.xlsx"

        saved = app.save_week_picks(1, {"Bills vs Jets": "BUF"}, workbook_path)

        self.assertFalse(saved)
        self.assertFalse(workbook_path.exists())

    def test_save_and_grade_picks_with_canonical_matchups(self) -> None:
        workbook_path = self.root / "nested" / "games.xlsx"

        saved = app.save_week_picks(1, {" NY Jets @ BUF ": "NY Jets"}, workbook_path)

        self.assertTrue(saved)
        picks = app.load_saved_picks(workbook_path)
        self.assertEqual(picks.iloc[0]["matchup"], "New York Jets @ Buffalo Bills")
        self.assertEqual(picks.iloc[0]["pick"], "New York Jets")

        hist_df = pd.DataFrame(
            [
                {
                    "week": 1,
                    "team1": "NY Jets",
                    "team2": "BUF",
                    "score1": 21,
                    "score2": 17,
                    "status": "Final",
                }
            ]
        )
        results = app.build_actual_results_by_week(hist_df)
        graded = app.grade_picks(picks, results)

        self.assertEqual(graded.iloc[0]["status"], "correct")
        self.assertEqual(graded.iloc[0]["result"], "✅ Correct")

    def test_build_actual_results_skips_invalid_teams(self) -> None:
        hist_df = pd.DataFrame(
            [
                {"week": 1, "team1": "Bad Team", "team2": "BUF", "score1": 10, "score2": 14, "status": "Final"},
            ]
        )

        results = app.build_actual_results_by_week(hist_df)

        self.assertTrue(results.empty)

    def test_build_actual_results_keeps_incomplete_final_pending(self) -> None:
        hist_df = pd.DataFrame(
            [
                {"week": 2, "team1": "BUF", "team2": "NYJ", "score1": 24, "score2": None, "status": "Final"},
            ]
        )

        results = app.build_actual_results_by_week(hist_df)

        self.assertEqual(results.iloc[0]["matchup"], "Buffalo Bills @ New York Jets")
        self.assertFalse(results.iloc[0]["is_final"])
        self.assertIsNone(results.iloc[0]["winner"])

    def test_build_actual_results_handles_live_and_postponed_statuses(self) -> None:
        hist_df = pd.DataFrame(
            [
                {"week": 3, "team1": "BUF", "team2": "NYJ", "score1": 14, "score2": 10, "status": "Live"},
                {"week": 3, "team1": "MIA", "team2": "NE", "score1": 7, "score2": 7, "status": "Postponed"},
            ]
        )

        results = app.build_actual_results_by_week(hist_df)

        self.assertEqual(results["is_final"].tolist(), [False, False])
        self.assertEqual(results["winner"].tolist(), [None, None])

    def test_grade_picks_marks_ties_as_pushes(self) -> None:
        picks = pd.DataFrame(
            [
                {
                    "week": 4,
                    "matchup": "BUF @ NY Jets",
                    "pick": "BUF",
                    "timestamp": "2026-09-17T00:00:00",
                }
            ]
        )
        hist_df = pd.DataFrame(
            [
                {"week": 4, "team1": "BUF", "team2": "NY Jets", "score1": 20, "score2": 20, "status": "Completed"},
            ]
        )

        results = app.build_actual_results_by_week(hist_df)
        graded = app.grade_picks(picks, results)

        self.assertEqual(graded.iloc[0]["winner"], "TIE")
        self.assertEqual(graded.iloc[0]["status"], "tie/push")
        self.assertEqual(graded.iloc[0]["result"], "➖ Tie/Push")

    def test_atomic_workbook_save_preserves_original_file_on_failure(self) -> None:
        workbook_path = self.root / "games.xlsx"
        original = Workbook()
        original.active.title = PICKS_SHEET
        original.active["A1"] = "original"
        original.save(workbook_path)
        original_bytes = workbook_path.read_bytes()
        before_names = sorted(path.name for path in self.root.iterdir())

        replacement = Workbook()
        replacement.active.title = PICKS_SHEET
        replacement.active["A1"] = "replacement"

        with patch.object(replacement, "save", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                app._atomic_workbook_save(replacement, workbook_path)

        self.assertEqual(workbook_path.read_bytes(), original_bytes)
        self.assertEqual(sorted(path.name for path in self.root.iterdir()), before_names)


if __name__ == "__main__":
    unittest.main()
